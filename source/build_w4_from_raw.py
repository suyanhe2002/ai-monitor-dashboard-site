from __future__ import annotations

import math
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


RAW_PATH = Path("/Users/lilblackmac/Desktop/质检记录列表6.22-6.25_合并.xlsx")
INVALID_PATH = Path("/Users/lilblackmac/Desktop/无效课节.xlsx")
REFERENCE_PATH = Path("/Users/lilblackmac/Desktop/AI监课&北极星/海外监课指标6月W4.xlsx")
OUTPUT_DIR = Path("/Users/lilblackmac/Documents/New project/outputs/w4_rebuild")
STANDARD_PATH = Path("/Users/lilblackmac/Desktop/AI监课&北极星/海外AI监课质检标准_优化版.xlsx")

PERIOD_LABEL = "W4"


RESULT_COLUMNS = {
    "背景布": "AI评测-背景布-结果",
    "工服": "AI评测-工服-结果",
    "光线": "AI评测-光线镜头-结果",
    "妆发": "AI评测-妆发-结果",
    "坐姿": "AI评测-坐姿-结果",
    "行为举止": "AI评测-行为举止-结果",
    "感染力": "AI评测-热情洋溢感染力强-结果",
    "课前作业提醒": "AI评测-学习习惯检查-结果",
    "课后作业布置": "AI评测-布置作业-结果",
    "知识点讲错": "AI评测-知识点讲错-结果",
    "知识点漏讲": "AI评测-知识点漏讲-结果",
    "课堂规则提及": "AI评测-课堂规则-结果",
    "读题审题": "AI评测-读题审题-结果",
    "关键提问": "AI评测-关键提问（依据详案）-结果",
    "情绪策略": "AI评测-使用情绪策略（依据详案）-结果",
    "课程预告": "AI评测-下节预告（动画+知识点）-结果",
}

DETAIL_COLUMNS = {
    "感染力分数": "AI评测-热情洋溢感染力强-明细",
    "声音洪亮分数": "AI评测-声音洪亮活力四射-明细",
    "开口时长(s)": "AI评测-学员开口时长（均值）-明细",
    "课中正确率": "AI评测-课堂正确率（均值）-明细",
    "上台平均数": "AI评测-学员参与度（均值）-明细",
    "主动性": "AI评测-主动性（均值 ）-明细",
    "学员数": "AI评测-学员参与度（学生占比）-明细",
    "镜头感": "AI评测-镜头感（均值）-明细",
    "愉悦度": "AI评测-愉悦度（均值）-明细",
    "有效互动次数": "AI评测-有效学员互动-明细",
}

TEACHER_OUTPUT_COLUMNS = [
    "老师",
    "业务部门",
    "小组",
    "样本课节数",
    "有效互动频次",
    "感染力合格率",
    "仪容仪表合格率",
    "背景布合格率",
    "工服合格率",
    "光线合格率",
    "妆发合格率",
    "坐姿合格率",
    "行为合格率",
    "课前作业提醒合格率",
    "课后作业布置合格率",
    "感染力分数",
    "声音洪亮分数",
    "知识点讲错通过率",
    "知识点漏讲通过率",
    "课堂规则提及通过率",
    "开口时长（s）",
    "课中正确率",
    "学员正面动作次数",
    "读题审题通过率",
    "关键提问通过率",
    "情绪策略使用通过率",
    "课程预告通过率",
]

SD_GROUP_COLUMNS = [
    "背景布合格率",
    "工服合格率",
    "光线合格率",
    "妆发合格率",
    "坐姿合格率",
    "行为合格率",
    "感染力合格率",
    "有效互动频次",
    "课前作业提醒合格率",
    "课后作业布置合格率",
    "感染力分数",
    "声音洪亮分数",
    "知识点讲错通过率",
    "知识点漏讲通过率",
    "课堂规则提及通过率",
    "开口时长（s）",
    "课中正确率",
    "学员正面动作次数",
    "读题审题通过率",
    "关键提问通过率",
    "情绪策略使用通过率",
    "课程预告通过率",
]

SUMMARY_SOURCE_MAP = {
    "背景布合格率": "背景布",
    "工服合格率": "工服",
    "光线合格率": "光线",
    "妆发合格率": "妆发",
    "坐姿合格率": "坐姿",
    "行为合格率": "行为举止",
    "感染力合格率": "感染力",
    "有效互动频次": "有效互动频次",
    "课前作业提醒合格率": "课前作业提醒",
    "课后作业布置合格率": "课后作业布置",
    "感染力分数": "感染力分数",
    "声音洪亮分数": "声音洪亮分数",
    "知识点讲错通过率": "知识点讲错",
    "知识点漏讲通过率": "知识点漏讲",
    "课堂规则提及通过率": "课堂规则提及",
    "开口时长（s）": "开口时长(s)",
    "课中正确率": "课中正确率",
    "学员正面动作次数": "主动性",
    "读题审题通过率": "读题审题",
    "关键提问通过率": "关键提问",
    "情绪策略使用通过率": "情绪策略",
    "课程预告通过率": "课程预告",
}

BUSINESS_DEPARTMENT_MAP = {
    "海外益智英语教学区": "英语",
    "海外益智粤语教学区": "粤语",
    "海外益智台湾教学区": "台湾",
    "海外益智外教教学区": "外教",
}

MONTH_SHEET_COLUMNS = [
    ("老师属性", None),
    ("小组", None),
    ("老师数", None),
    ("上台次数（BI）", "target"),
    ("上台次数（BI）", "actual"),
    ("上台次数（BI）", "gap"),
    ("感染力合格率", "target"),
    ("感染力合格率", "actual"),
    ("感染力合格率", "gap"),
    ("有效互动频次", "target"),
    ("有效互动频次", "actual"),
    ("有效互动频次", "gap"),
    ("仪容仪表合格率", "actual"),
    ("背景布合格率", "actual"),
    ("工服合格率", "actual"),
    ("光线合格率", "actual"),
    ("妆发合格率", "actual"),
    ("坐姿合格率", "actual"),
    ("行为合格率", "actual"),
    ("课前作业提醒合格率", "actual"),
    ("课后作业布置合格率", "actual"),
    ("感染力分数", "actual"),
    ("声音洪亮分数", "actual"),
    ("知识点讲错通过率", "actual"),
    ("知识点漏讲通过率", "actual"),
    ("课堂规则提及通过率", "actual"),
    ("开口时长（s）", "actual"),
    ("课中正确率", "actual"),
    ("学员正面动作次数", "actual"),
    ("读题审题通过率", "actual"),
    ("关键提问通过率", "actual"),
    ("情绪策略使用通过率", "actual"),
    ("课程预告通过率", "actual"),
]

MONTH_TARGETS = {
    "上台次数（BI）": 6.0,
    "感染力合格率": 0.5,
}


def safe_float(value: object) -> float | None:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def parse_pct_from_text(text: object) -> float | None:
    if text is None or (isinstance(text, float) and math.isnan(text)):
        return None
    match = re.search(r"([0-9]+(?:\.[0-9]+)?)%", str(text))
    return None if not match else float(match.group(1)) / 100


def parse_number_after_colon(text: object) -> float | None:
    if text is None or (isinstance(text, float) and math.isnan(text)):
        return None
    match = re.search(r":\s*([0-9]+(?:\.[0-9]+)?)", str(text))
    return None if not match else float(match.group(1))


def parse_active_student_count(text: object) -> float | None:
    if text is None or (isinstance(text, float) and math.isnan(text)):
        return None
    match = re.search(r"上下台≥1的学生数[:：]\s*([0-9]+)", str(text))
    if match:
        return float(match.group(1))
    match = re.search(r"学生数[:：]\s*([0-9]+)", str(text))
    if match:
        return float(match.group(1))
    match = re.search(r"掌握学生数[:：]\s*([0-9]+)", str(text))
    return None if not match else float(match.group(1))


def parse_total_students_from_participation(text: object) -> float | None:
    if text is None or (isinstance(text, float) and math.isnan(text)):
        return None
    raw = str(text)
    count_match = re.search(r"上下台≥1的学生数[:：]\s*([0-9]+)", raw)
    ratio_match = re.search(r"上下台≥1的学生占比[:：]\s*([0-9]+(?:\.[0-9]+)?)%", raw)
    if not count_match or not ratio_match:
        return None
    count = float(count_match.group(1))
    ratio = float(ratio_match.group(1)) / 100
    if ratio <= 0:
        return None
    return round(count / ratio, 0)


def parse_interaction_frequency(text: object) -> float | None:
    if text is None or (isinstance(text, float) and math.isnan(text)):
        return None
    total_match = re.search(r"老师互动动作总数[:：]\s*([0-9]+(?:\.[0-9]+)?)", str(text))
    return None if not total_match else float(total_match.group(1))


def teacher_alias(name: object) -> str:
    if name is None:
        return ""
    text = str(name).strip()
    return re.sub(r"[（(].*?[）)]", "", text).strip()


def extract_department_parts(value: object) -> tuple[str, str]:
    if value is None:
        return "", ""
    parts = [part.strip() for part in str(value).split("/") if part and str(part).strip()]
    if len(parts) >= 2:
        return parts[-2], parts[-1]
    if len(parts) == 1:
        return parts[0], ""
    return "", ""


def load_invalid_lessons() -> set[int]:
    raw = pd.read_excel(INVALID_PATH, header=None)
    header_row = 7
    data = raw.iloc[header_row:].copy()
    data.columns = raw.iloc[header_row].tolist()
    data = data[data["课节ID"].notna()].copy()
    ids = pd.to_numeric(data["课节ID"], errors="coerce").dropna().astype(int)
    return set(ids.tolist())


def build_detail() -> pd.DataFrame:
    df = pd.read_excel(RAW_PATH)
    invalid_ids = load_invalid_lessons()

    detail = pd.DataFrame()
    detail["周期"] = PERIOD_LABEL
    detail["样本数"] = (~df["房间号"].isin(invalid_ids)).astype(int)
    detail["老师"] = df["教师姓名"].map(teacher_alias)
    detail["老师工号"] = df["教师工号"].astype(str).str.strip()
    detail["业务部门"], detail["小组"] = zip(*df["教师部门"].map(extract_department_parts))
    detail["课节ID"] = pd.to_numeric(df["房间号"], errors="coerce")
    detail["日期"] = pd.to_datetime(df["上课时间"].astype(str).str.extract(r"(^\d{4}-\d{2}-\d{2})")[0], errors="coerce")

    for output_name, source_col in RESULT_COLUMNS.items():
        detail[output_name] = (df[source_col] == "通过").astype(float)

    detail["感染力分数"] = df[DETAIL_COLUMNS["感染力分数"]].map(parse_pct_from_text)
    detail["声音洪亮分数"] = df[DETAIL_COLUMNS["声音洪亮分数"]].map(parse_pct_from_text)
    detail["开口时长(s)"] = df[DETAIL_COLUMNS["开口时长(s)"]].map(parse_number_after_colon)
    detail["课中正确率"] = df[DETAIL_COLUMNS["课中正确率"]].map(parse_pct_from_text)
    detail["上台平均数"] = df[DETAIL_COLUMNS["上台平均数"]].map(parse_number_after_colon)
    detail["主动性"] = df[DETAIL_COLUMNS["主动性"]].map(parse_number_after_colon)
    detail["上台学生数"] = df[DETAIL_COLUMNS["学员数"]].map(parse_active_student_count)
    detail["学员数"] = df[DETAIL_COLUMNS["学员数"]].map(parse_total_students_from_participation)
    detail["镜头感"] = df[DETAIL_COLUMNS["镜头感"]].map(parse_pct_from_text)
    detail["愉悦度"] = df[DETAIL_COLUMNS["愉悦度"]].map(parse_pct_from_text)
    detail["有效互动次数"] = df[DETAIL_COLUMNS["有效互动次数"]].map(parse_interaction_frequency)
    detail["有效互动频次"] = detail.apply(
        lambda row: None
        if row["学员数"] is None or pd.isna(row["学员数"]) or row["学员数"] <= 0
        else safe_float(row["有效互动次数"] / row["学员数"]) if row["有效互动次数"] is not None and not pd.isna(row["有效互动次数"]) else None,
        axis=1,
    )
    return detail


def ratio_metric(group: pd.DataFrame, col: str) -> float | None:
    valid = group[group["样本数"] == 1]
    if valid.empty:
        return None
    return safe_float(valid[col].sum() / len(valid))


def average_metric(group: pd.DataFrame, col: str) -> float | None:
    valid = group[col].dropna()
    if valid.empty:
        return None
    return safe_float(valid.mean())


def average_metric_valid_only(group: pd.DataFrame, col: str) -> float | None:
    filtered = group[group["样本数"] == 1][col].dropna()
    if filtered.empty:
        return None
    return safe_float(filtered.mean())


def build_teacher_summary(detail: pd.DataFrame) -> pd.DataFrame:
    records: list[dict[str, object]] = []
    for (teacher, dept, group_name), group in detail.groupby(["老师", "业务部门", "小组"], dropna=False):
        row: dict[str, object] = {
            "老师": teacher,
            "业务部门": dept,
            "小组": group_name,
            "样本课节数": int(group["课节ID"].count()),
            "有效互动频次": average_metric_valid_only(group, "有效互动频次"),
            "感染力合格率": ratio_metric(group, "感染力"),
            "背景布合格率": ratio_metric(group, "背景布"),
            "工服合格率": ratio_metric(group, "工服"),
            "光线合格率": ratio_metric(group, "光线"),
            "妆发合格率": ratio_metric(group, "妆发"),
            "坐姿合格率": ratio_metric(group, "坐姿"),
            "行为合格率": ratio_metric(group, "行为举止"),
            "课前作业提醒合格率": ratio_metric(group, "课前作业提醒"),
            "课后作业布置合格率": ratio_metric(group, "课后作业布置"),
            "感染力分数": average_metric(group, "感染力分数"),
            "声音洪亮分数": average_metric(group, "声音洪亮分数"),
            "知识点讲错通过率": ratio_metric(group, "知识点讲错"),
            "知识点漏讲通过率": ratio_metric(group, "知识点漏讲"),
            "课堂规则提及通过率": ratio_metric(group, "课堂规则提及"),
            "开口时长（s）": average_metric_valid_only(group, "开口时长(s)"),
            "课中正确率": average_metric_valid_only(group, "课中正确率"),
            "学员正面动作次数": average_metric_valid_only(group, "主动性"),
            "读题审题通过率": ratio_metric(group, "读题审题"),
            "关键提问通过率": ratio_metric(group, "关键提问"),
            "情绪策略使用通过率": ratio_metric(group, "情绪策略"),
            "课程预告通过率": ratio_metric(group, "课程预告"),
        }
        grooming_items = [
            row["背景布合格率"],
            row["工服合格率"],
            row["光线合格率"],
            row["妆发合格率"],
            row["坐姿合格率"],
            row["行为合格率"],
        ]
        grooming_values = [value for value in grooming_items if value is not None]
        row["仪容仪表合格率"] = safe_float(sum(grooming_values) / len(grooming_values)) if grooming_values else None
        records.append(row)

    teacher_df = pd.DataFrame(records)
    teacher_df = teacher_df.sort_values(["业务部门", "小组", "老师"], kind="stable").reset_index(drop=True)
    return teacher_df[TEACHER_OUTPUT_COLUMNS]


def build_reference_comparison(summary: pd.DataFrame) -> pd.DataFrame:
    ref = pd.read_excel(REFERENCE_PATH, sheet_name="W4老师达成情况", header=1)
    ref = ref[ref["昵称"].notna()].copy()
    ref["老师"] = ref["昵称"].astype(str).str.strip()
    ref_comp = ref.rename(
        columns={
            "大组": "业务部门",
            "背景布\n合格率": "背景布合格率",
            "工服\n合格率": "工服合格率",
            "光线\n合格率": "光线合格率",
            "妆发\n合格率": "妆发合格率",
            "坐姿\n合格率": "坐姿合格率",
            "行为\n合格率": "行为合格率",
            "感染力\n分数": "感染力分数",
            "开口时长\n（s）": "开口时长（s）",
            "学员正面动作次数": "学员正面动作次数",
            "情绪策略使用通过率": "情绪策略使用通过率",
        }
    )
    compare_cols = [
        "老师",
        "业务部门",
        "小组",
        "有效互动频次",
        "感染力合格率",
        "仪容仪表合格率",
        "背景布合格率",
        "工服合格率",
        "光线合格率",
        "妆发合格率",
        "坐姿合格率",
        "行为合格率",
        "课前作业提醒合格率",
        "课后作业布置合格率",
        "感染力分数",
        "声音洪亮分数",
        "知识点讲错通过率",
        "知识点漏讲通过率",
        "课堂规则提及通过率",
        "开口时长（s）",
        "课中正确率",
        "学员正面动作次数",
        "读题审题通过率",
        "关键提问通过率",
        "情绪策略使用通过率",
        "课程预告通过率",
    ]
    merged = summary.merge(ref_comp[compare_cols], on=["老师", "业务部门", "小组"], how="left", suffixes=("_calc", "_ref"))
    for col in compare_cols[3:]:
        merged[f"{col}_diff"] = merged[f"{col}_calc"] - merged[f"{col}_ref"]
    return merged


def build_team_and_group_summary(detail: pd.DataFrame, teacher_summary: pd.DataFrame) -> pd.DataFrame:
    valid_detail = detail[detail["样本数"] == 1].copy()

    rows: list[dict[str, object]] = []

    ratio_metrics = {
        "背景布合格率",
        "工服合格率",
        "光线合格率",
        "妆发合格率",
        "坐姿合格率",
        "行为合格率",
        "感染力合格率",
        "课前作业提醒合格率",
        "课后作业布置合格率",
        "知识点讲错通过率",
        "知识点漏讲通过率",
        "课堂规则提及通过率",
        "课中正确率",
        "读题审题通过率",
        "关键提问通过率",
        "情绪策略使用通过率",
        "课程预告通过率",
    }

    for dept, group in valid_detail.groupby("业务部门", dropna=False):
        teacher_count = int(
            teacher_summary.loc[teacher_summary["业务部门"] == dept, "老师"].nunique()
        )
        row = {
            "层级": "团队",
            "名称": dept,
            "业务部门": dept,
            "样本数": int(len(group)),
            "老师数": teacher_count,
        }
        for col in SD_GROUP_COLUMNS:
            source_col = SUMMARY_SOURCE_MAP[col]
            values = pd.to_numeric(group[source_col], errors="coerce")
            row[col] = safe_float(values.mean()) if not values.dropna().empty else None
        rows.append(row)

    for (dept, group_name), group in valid_detail.groupby(["业务部门", "小组"], dropna=False):
        teacher_count = int(
            teacher_summary.loc[
                (teacher_summary["业务部门"] == dept) & (teacher_summary["小组"] == group_name),
                "老师",
            ].nunique()
        )
        row = {
            "层级": "小组",
            "名称": group_name,
            "业务部门": dept,
            "样本数": int(len(group)),
            "老师数": teacher_count,
        }
        for col in SD_GROUP_COLUMNS:
            source_col = SUMMARY_SOURCE_MAP[col]
            values = pd.to_numeric(group[source_col], errors="coerce")
            row[col] = safe_float(values.mean()) if not values.dropna().empty else None
        rows.append(row)

    result = pd.DataFrame(rows)
    team_order = ["海外益智英语教学区", "海外益智粤语教学区", "海外益智台湾教学区", "海外益智外教教学区"]
    result["层级排序"] = result["层级"].map({"团队": 0, "小组": 1}).fillna(9)
    result["团队排序"] = result["业务部门"].fillna(result["名称"]).map({name: i for i, name in enumerate(team_order)}).fillna(99)
    result = result.sort_values(["层级排序", "团队排序", "名称"], kind="stable").drop(columns=["层级排序", "团队排序"])
    return result


def load_standard_targets() -> dict[str, float | None]:
    df = pd.read_excel(STANDARD_PATH)
    metric_map = {
        "背景布": "背景布合格率",
        "工服": "工服合格率",
        "光线镜头": "光线合格率",
        "妆发": "妆发合格率",
        "坐姿": "坐姿合格率",
        "行为举止": "行为合格率",
        "热情洋溢感染力强": "感染力合格率",
        "知识点讲错": "知识点讲错通过率",
        "知识点漏讲": "知识点漏讲通过率",
        "课堂规则": "课堂规则提及通过率",
        "学习习惯检查": "课前作业提醒合格率",
        "学员开口时长（均值）": "开口时长（s）",
        "课堂正确率（均值）": "课中正确率",
        "学员参与度（均值）": "上台次数（BI）",
        "主动性（均值 ）": "学员正面动作次数",
        "布置作业": "课后作业布置合格率",
        "下节预告（动画+知识点）": "课程预告通过率",
    }
    targets: dict[str, float | None] = {}
    for _, row in df.iterrows():
        indicator = str(row.get("指标", "")).strip()
        target_col = metric_map.get(indicator)
        if not target_col:
            continue
        text = str(row.get("质检标准调整", "") or "").replace(" ", "")
        pct_match = re.search(r"A：.*?([0-9]+(?:\.[0-9]+)?)%", text)
        num_match = re.search(r"A：.*?([0-9]+(?:\.[0-9]+)?)", text)
        target = None
        if pct_match:
            target = float(pct_match.group(1)) / 100
        elif num_match:
            target = float(num_match.group(1))
        if "分钟" in text and target is not None:
            target *= 60
        targets[target_col] = target
    targets.update(MONTH_TARGETS)
    return targets


def build_month_achievement(teacher_summary: pd.DataFrame) -> pd.DataFrame:
    business = teacher_summary[teacher_summary["业务部门"].isin(BUSINESS_DEPARTMENT_MAP.keys())].copy()
    targets = load_standard_targets()
    rows: list[dict[str, object]] = []
    ordered_departments = list(BUSINESS_DEPARTMENT_MAP.keys())

    for dept in ordered_departments:
        dept_df = business[business["业务部门"] == dept].copy()
        if dept_df.empty:
            continue
        dept_label = BUSINESS_DEPARTMENT_MAP[dept]
        for group_name, group_df in dept_df.groupby("小组", dropna=False):
            row: dict[str, object] = {
                "老师属性": dept_label,
                "小组": group_name,
                "老师数": int(group_df["老师"].nunique()),
            }
            for col in TEACHER_OUTPUT_COLUMNS[4:]:
                row[col] = safe_float(pd.to_numeric(group_df[col], errors="coerce").mean())
            row["上台次数（BI）"] = row.get("上台次数（BI）")
            rows.append(row)

        total_row = {
            "老师属性": dept_label,
            "小组": "",
            "老师数": int(dept_df["老师"].nunique()),
        }
        for col in TEACHER_OUTPUT_COLUMNS[4:]:
            total_row[col] = safe_float(pd.to_numeric(dept_df[col], errors="coerce").mean())
        rows.append(total_row)

    month_df = pd.DataFrame(rows)
    for metric, target in targets.items():
        month_df[f"{metric}_target"] = target
        if metric in month_df.columns and target is not None:
            month_df[f"{metric}_gap"] = pd.to_numeric(month_df[metric], errors="coerce") - target
        else:
            month_df[f"{metric}_gap"] = None
    return month_df


def export_workbook(detail: pd.DataFrame, summary: pd.DataFrame, team_group_summary: pd.DataFrame, month_df: pd.DataFrame) -> Path:
    output_path = OUTPUT_DIR / "海外监课指标6月W4_重建版.xlsx"
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        detail.to_excel(writer, sheet_name="质检明细", index=False)
        summary.to_excel(writer, sheet_name="W4老师达成情况", index=False, startrow=1)
        team_group_summary.to_excel(writer, sheet_name="数透", index=False, startrow=2)

        sheet_rows = []
        title_row = ["6.22-6.25"] + [""] * (len(MONTH_SHEET_COLUMNS) - 1)
        top_header = []
        sub_header = []
        for col, sub in MONTH_SHEET_COLUMNS:
            top_header.append(col)
            if sub is None:
                sub_header.append("")
            elif sub == "target":
                sub_header.append("目标值")
            elif sub == "actual":
                sub_header.append("实际值")
            else:
                sub_header.append("GAP")
        sheet_rows.append(title_row)
        sheet_rows.append(top_header)
        sheet_rows.append(sub_header)
        for _, row in month_df.iterrows():
            values = []
            for col, sub in MONTH_SHEET_COLUMNS:
                if sub is None:
                    values.append(row.get(col))
                elif sub == "target":
                    values.append(row.get(f"{col}_target"))
                elif sub == "actual":
                    values.append(row.get(col))
                else:
                    values.append(row.get(f"{col}_gap"))
            sheet_rows.append(values)
        pd.DataFrame(sheet_rows).to_excel(writer, sheet_name="6月达成情况", index=False, header=False)

    wb = load_workbook(output_path)
    thin_gray = Side(style="thin", color="D9D9D9")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    sub_fill = PatternFill("solid", fgColor="FBE5D6")
    title_fill = PatternFill("solid", fgColor="1F4E78")
    yellow_fill = PatternFill("solid", fgColor="FFF2CC")
    bold_font = Font(name="Microsoft YaHei", bold=True)
    white_bold_font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei")

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2" if ws.title != "6月达成情况" else "A4"
        for row in ws.iter_rows():
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
                cell.border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    detail_ws = wb["质检明细"]
    for cell in detail_ws[1]:
        cell.fill = yellow_fill
        cell.font = bold_font
    for col in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L"]:
        detail_ws.column_dimensions[col].width = 14
    for col in ["M", "N", "O", "P", "Q", "R", "S", "T", "U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC", "AD"]:
        detail_ws.column_dimensions[col].width = 16

    teacher_ws = wb["W4老师达成情况"]
    teacher_ws["A1"] = "目标值"
    targets = load_standard_targets()
    for idx, col_name in enumerate(summary.columns, start=1):
        teacher_ws.cell(row=2, column=idx).fill = header_fill
        teacher_ws.cell(row=2, column=idx).font = bold_font
        teacher_ws.cell(row=1, column=idx).fill = sub_fill
        teacher_ws.cell(row=1, column=idx).font = bold_font
        teacher_ws.cell(row=1, column=idx).value = targets.get(col_name) if col_name in targets else "-"
        teacher_ws.column_dimensions[teacher_ws.cell(row=2, column=idx).column_letter].width = 14

    group_ws = wb["数透"]
    group_ws["A1"] = "周期"
    group_ws["B1"] = PERIOD_LABEL
    for row_idx in [3]:
        for cell in group_ws[row_idx]:
            cell.fill = header_fill
            cell.font = bold_font
    for col in range(1, group_ws.max_column + 1):
        group_ws.column_dimensions[group_ws.cell(row=3, column=col).column_letter].width = 16

    month_ws = wb["6月达成情况"]
    month_ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=month_ws.max_column)
    month_ws["A1"].fill = title_fill
    month_ws["A1"].font = white_bold_font
    month_ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    for row_idx in [2, 3]:
        fill = header_fill if row_idx == 2 else sub_fill
        for cell in month_ws[row_idx]:
            cell.fill = fill
            cell.font = bold_font
    for col in range(1, month_ws.max_column + 1):
        letter = month_ws.cell(row=2, column=col).column_letter
        month_ws.column_dimensions[letter].width = 14
    month_ws.column_dimensions["A"].width = 10
    month_ws.column_dimensions["B"].width = 12
    month_ws.column_dimensions["C"].width = 8
    for row in month_ws.iter_rows(min_row=4, max_row=month_ws.max_row):
        if row[1].value in ("", None):
            for cell in row:
                cell.font = bold_font

    percent_keywords = ["合格率", "通过率", "正确率", "分数"]
    for ws_name in ["W4老师达成情况", "数透", "6月达成情况"]:
        ws = wb[ws_name]
        for row in ws.iter_rows():
            for cell in row:
                header_value = ws.cell(row=2 if ws_name == "6月达成情况" else (2 if ws_name == "W4老师达成情况" else 3), column=cell.column).value
                header_text = "" if header_value is None else str(header_value)
                if any(keyword in header_text for keyword in percent_keywords) and isinstance(cell.value, (int, float)):
                    cell.number_format = "0.0%"
                elif "GAP" in str(ws.cell(row=3 if ws_name == "6月达成情况" else 1, column=cell.column).value) and isinstance(cell.value, (int, float)):
                    if any(keyword in header_text for keyword in percent_keywords):
                        cell.number_format = "0.0%"
                    else:
                        cell.number_format = "0.0"
                elif isinstance(cell.value, (int, float)):
                    cell.number_format = "0.0"

    wb.save(output_path)
    return output_path


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    detail = build_detail()
    summary = build_teacher_summary(detail)
    comparison = build_reference_comparison(summary)
    team_group_summary = build_team_and_group_summary(detail, summary)
    month_df = build_month_achievement(summary)

    detail.to_csv(OUTPUT_DIR / "w4_detail_rebuilt.csv", index=False, encoding="utf-8-sig")
    summary.to_csv(OUTPUT_DIR / "w4_teacher_summary_rebuilt.csv", index=False, encoding="utf-8-sig")
    comparison.to_csv(OUTPUT_DIR / "w4_teacher_summary_compare.csv", index=False, encoding="utf-8-sig")
    team_group_summary.to_csv(OUTPUT_DIR / "w4_group_summary_rebuilt.csv", index=False, encoding="utf-8-sig")
    month_df.to_csv(OUTPUT_DIR / "w4_month_achievement_rebuilt.csv", index=False, encoding="utf-8-sig")
    workbook_path = export_workbook(detail, summary, team_group_summary, month_df)

    print(OUTPUT_DIR / "w4_detail_rebuilt.csv")
    print(OUTPUT_DIR / "w4_teacher_summary_rebuilt.csv")
    print(OUTPUT_DIR / "w4_teacher_summary_compare.csv")
    print(OUTPUT_DIR / "w4_group_summary_rebuilt.csv")
    print(OUTPUT_DIR / "w4_month_achievement_rebuilt.csv")
    print(workbook_path)


if __name__ == "__main__":
    main()
