from __future__ import annotations

from pathlib import Path
import sys

import pandas as pd
from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import build_7w1_from_raw as src


TARGET_PATH = Path("/Users/lilblackmac/Desktop/海外监课指标7月W1_模板回填版.xlsx")

OVERRIDE_GROUPS = {
    "督导": {"Jessica老师", "Martin老师", "Peggy老師", "YOYO老師"},
    "职能岗": {"Dainty老师", "小加老師", "coco老師", "彩笔老师"},
}


TOP_ORDER = [
    "英语",
    "英语01组",
    "英语02组",
    "英语03组",
    "英语04组",
    "英语06组",
    "英语07组",
    "粤语",
    "粤语01组",
    "粤语02组",
    "粤语03组",
    "粤语04组",
    "粤语05组",
    "粤语06组",
    "台教",
    "菲教",
    "督导",
    "职能岗",
    "总计",
]

LEFT_ORDER = [
    "英语",
    "英语01组",
    "英语02组",
    "英语03组",
    "英语04组",
    "英语06组",
    "英语07组",
    "粤语",
    "粤语01组",
    "粤语02组",
    "粤语03组",
    "粤语04组",
    "粤语05组",
    "粤语06组",
]

RIGHT_ORDER = TOP_ORDER


def dept_label(dept: str) -> str:
    if dept in src.LABEL_MAP:
        return src.LABEL_MAP[dept]
    dept_str = str(dept)
    if "质量" in dept_str or "运营" in dept_str:
        return "督导"
    if "教学部" in dept_str:
        return "职能岗"
    return dept_str


def build_top_map(teacher_summary: pd.DataFrame, bi: pd.DataFrame) -> dict[str, list[object]]:
    rows: dict[str, list[object]] = {}
    for dept in teacher_summary["大组"].dropna().drop_duplicates():
        label = dept_label(dept)
        tsub = teacher_summary[teacher_summary["大组"] == dept].copy()
        bsub = bi[bi["大组"] == dept].copy()
        rows[label] = [
            label,
            src.safe_float(pd.to_numeric(bsub["上台次数（BI）"], errors="coerce").mean()),
            int(tsub["昵称"].nunique()),
        ]
        for group_name, g in tsub.groupby("小组", dropna=False):
            group_label = "兼职" if str(group_name).strip() == "兼职老师组" else ("" if pd.isna(group_name) else str(group_name))
            gb = bi[(bi["大组"] == dept) & (bi["小组"] == group_name)].copy()
            rows[group_label] = [
                group_label,
                src.safe_float(pd.to_numeric(gb["上台次数（BI）"], errors="coerce").mean()),
                int(g["昵称"].nunique()),
            ]
    rows["总计"] = [
        "总计",
        src.safe_float(pd.to_numeric(teacher_summary["上台次数（BI）"], errors="coerce").mean()),
        int(teacher_summary["昵称"].nunique()),
    ]
    for label, names in OVERRIDE_GROUPS.items():
        tsub = teacher_summary[teacher_summary["昵称"].isin(names)].copy()
        rows[label] = [
            label,
            src.safe_float(pd.to_numeric(tsub["上台次数（BI）"], errors="coerce").mean()),
            int(tsub["昵称"].nunique()),
        ]
    return rows


def build_left_map(teacher_summary: pd.DataFrame, detail: pd.DataFrame) -> dict[str, list[object]]:
    valid_detail = detail[detail["样本数"] == 1].copy()
    rows: dict[str, list[object]] = {}
    seen: set[tuple[str, object]] = set()
    order_rows: list[tuple[str, object]] = []

    for _, row in teacher_summary[["大组", "小组"]].drop_duplicates().iterrows():
        dept = row["大组"]
        group = row["小组"]
        if (dept, None) not in seen:
            seen.add((dept, None))
            order_rows.append((dept, None))
        key = (dept, group)
        if key not in seen:
            seen.add(key)
            order_rows.append(key)

    for dept, group in order_rows:
        if group is None:
            sub = teacher_summary[teacher_summary["大组"] == dept].copy()
            dsub = valid_detail[valid_detail["教师部门5"] == dept].copy()
            label = dept_label(dept)
        else:
            sub = teacher_summary[(teacher_summary["大组"] == dept) & (teacher_summary["小组"] == group)].copy()
            dsub = valid_detail[(valid_detail["教师部门5"] == dept) & (valid_detail["教师部门6"] == group)].copy()
            label = "兼职" if str(group).strip() == "兼职老师组" else str(group)
        if sub.empty:
            continue
        rows[label] = [
            label,
            int(len(dsub)),
            int(sub["昵称"].nunique()),
            src.safe_float(pd.to_numeric(sub["背景布\n合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["工服\n合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["光线\n合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["妆发\n合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["坐姿\n合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["行为\n合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["感染力合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["有效互动频次"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["课前作业提醒合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["课后作业布置合格率"], errors="coerce").mean()),
        ]

    rows["总计"] = [
        "总计",
        int(len(valid_detail)),
        int(teacher_summary["昵称"].nunique()),
        src.safe_float(pd.to_numeric(teacher_summary["背景布\n合格率"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["工服\n合格率"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["光线\n合格率"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["妆发\n合格率"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["坐姿\n合格率"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["行为\n合格率"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["感染力合格率"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["有效互动频次"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["课前作业提醒合格率"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["课后作业布置合格率"], errors="coerce").mean()),
    ]
    for label, names in OVERRIDE_GROUPS.items():
        sub = teacher_summary[teacher_summary["昵称"].isin(names)].copy()
        dsub = valid_detail[valid_detail["教师艺名"].isin(names)].copy()
        rows[label] = [
            label,
            int(len(dsub)),
            int(sub["昵称"].nunique()),
            src.safe_float(pd.to_numeric(sub["背景布\n合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["工服\n合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["光线\n合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["妆发\n合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["坐姿\n合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["行为\n合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["感染力合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["有效互动频次"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["课前作业提醒合格率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(sub["课后作业布置合格率"], errors="coerce").mean()),
        ]
    return rows


def build_right_map(teacher_summary: pd.DataFrame, detail: pd.DataFrame) -> dict[str, list[object]]:
    valid_detail = detail[detail["样本数"] == 1].copy()
    rows: dict[str, list[object]] = {}
    for dept in teacher_summary["大组"].dropna().drop_duplicates():
        label = dept_label(dept)
        tsub = teacher_summary[teacher_summary["大组"] == dept].copy()
        dsub = valid_detail[valid_detail["教师部门5"] == dept].copy()
        rows[label] = [
            label,
            int(len(dsub)),
            int(tsub["昵称"].nunique()),
            src.safe_float(pd.to_numeric(tsub["感染力\n分数"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["声音洪亮分数"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["知识点讲错通过率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["知识点漏讲通过率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["课堂规则提及通过率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["开口时长\n（s）"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["课中正确率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["上台次数（BI）"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["学员正面动作次数"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["读题审题通过率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["关键提问通过率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["情绪策略使用通过率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["课程预告通过率"], errors="coerce").mean()),
        ]
        for group_name, g in tsub.groupby("小组", dropna=False):
            group_label = "兼职" if str(group_name).strip() == "兼职老师组" else ("" if pd.isna(group_name) else str(group_name))
            gd = valid_detail[(valid_detail["教师部门5"] == dept) & (valid_detail["教师部门6"] == group_name)].copy()
            rows[group_label] = [
                group_label,
                int(len(gd)),
                int(g["昵称"].nunique()),
                src.safe_float(pd.to_numeric(g["感染力\n分数"], errors="coerce").mean()),
                src.safe_float(pd.to_numeric(g["声音洪亮分数"], errors="coerce").mean()),
                src.safe_float(pd.to_numeric(g["知识点讲错通过率"], errors="coerce").mean()),
                src.safe_float(pd.to_numeric(g["知识点漏讲通过率"], errors="coerce").mean()),
                src.safe_float(pd.to_numeric(g["课堂规则提及通过率"], errors="coerce").mean()),
                src.safe_float(pd.to_numeric(g["开口时长\n（s）"], errors="coerce").mean()),
                src.safe_float(pd.to_numeric(g["课中正确率"], errors="coerce").mean()),
                src.safe_float(pd.to_numeric(g["上台次数（BI）"], errors="coerce").mean()),
                src.safe_float(pd.to_numeric(g["学员正面动作次数"], errors="coerce").mean()),
                src.safe_float(pd.to_numeric(g["读题审题通过率"], errors="coerce").mean()),
                src.safe_float(pd.to_numeric(g["关键提问通过率"], errors="coerce").mean()),
                src.safe_float(pd.to_numeric(g["情绪策略使用通过率"], errors="coerce").mean()),
                src.safe_float(pd.to_numeric(g["课程预告通过率"], errors="coerce").mean()),
            ]

    rows["总计"] = [
        "总计",
        int(len(valid_detail)),
        int(teacher_summary["昵称"].nunique()),
        src.safe_float(pd.to_numeric(teacher_summary["感染力\n分数"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["声音洪亮分数"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["知识点讲错通过率"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["知识点漏讲通过率"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["课堂规则提及通过率"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["开口时长\n（s）"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["课中正确率"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["上台次数（BI）"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["学员正面动作次数"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["读题审题通过率"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["关键提问通过率"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["情绪策略使用通过率"], errors="coerce").mean()),
        src.safe_float(pd.to_numeric(teacher_summary["课程预告通过率"], errors="coerce").mean()),
    ]
    for label, names in OVERRIDE_GROUPS.items():
        tsub = teacher_summary[teacher_summary["昵称"].isin(names)].copy()
        dsub = valid_detail[valid_detail["教师艺名"].isin(names)].copy()
        rows[label] = [
            label,
            int(len(dsub)),
            int(tsub["昵称"].nunique()),
            src.safe_float(pd.to_numeric(tsub["感染力\n分数"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["声音洪亮分数"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["知识点讲错通过率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["知识点漏讲通过率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["课堂规则提及通过率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["开口时长\n（s）"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["课中正确率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["上台次数（BI）"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["学员正面动作次数"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["读题审题通过率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["关键提问通过率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["情绪策略使用通过率"], errors="coerce").mean()),
            src.safe_float(pd.to_numeric(tsub["课程预告通过率"], errors="coerce").mean()),
        ]
    return rows


def fill_row(ws, row_idx: int, start_col: int, values: list[object]) -> None:
    for offset, value in enumerate(values):
        ws.cell(row_idx, start_col + offset, value)


def main() -> None:
    raw = src.load_raw()
    detail = src.build_detail(raw)
    bi = src.load_bi()
    teacher_summary = src.build_teacher_summary(detail, bi)

    top_map = build_top_map(teacher_summary, bi)
    left_map = build_left_map(teacher_summary, detail)
    right_map = build_right_map(teacher_summary, detail)

    wb = load_workbook(TARGET_PATH)
    ws = wb["数透"]

    ws["B2"] = "W1"
    ws["R1"] = "周期"
    ws["S1"] = "W1"
    ws["A32"] = "周期"
    ws["B32"] = "W1"

    # Top block A4:G23
    ws["A4"] = "行标签"
    ws["B4"] = "上台次数 "
    ws["C4"] = "去重老师数"
    ws["E4"] = "行标签"
    ws["F4"] = "上台次数 "
    ws["G4"] = "去重老师数"
    for idx, label in enumerate(TOP_ORDER, start=5):
        row = top_map[label]
        fill_row(ws, idx, 1, row)
        fill_row(ws, idx, 5, row)
    for idx in range(24, 26):
        for col in range(1, 8):
            ws.cell(idx, col, None)

    # Weekly left block A34:M48
    ws["A34"] = "团队"
    ws["B34"] = "样本数 "
    ws["C34"] = "老师数"
    ws["D34"] = "背景布合格率"
    ws["E34"] = "工服合格率"
    ws["F34"] = "光线合格率"
    ws["G34"] = "妆发合格率"
    ws["H34"] = "坐姿合格率"
    ws["I34"] = "行为合格率"
    ws["J34"] = "感染力合格率"
    ws["K34"] = "有效互动频次"
    ws["L34"] = "课前作业提醒合格率"
    ws["M34"] = "课后作业布置合格率"
    for idx, label in enumerate(LEFT_ORDER, start=35):
        fill_row(ws, idx, 1, left_map[label])
    for idx in range(49, 59):
        for col in range(1, 14):
            ws.cell(idx, col, None)

    # Monthly/period left block A59:M81
    ws["A59"] = "团队"
    ws["B59"] = "样本数 "
    ws["C59"] = "老师数"
    ws["D59"] = "背景布合格率"
    ws["E59"] = "工服合格率"
    ws["F59"] = "光线合格率"
    ws["G59"] = "妆发合格率"
    ws["H59"] = "坐姿合格率"
    ws["I59"] = "行为合格率"
    ws["J59"] = "感染力合格率"
    ws["K59"] = "有效互动频次"
    ws["L59"] = "课前作业提醒合格率"
    ws["M59"] = "课后作业布置合格率"
    for idx, label in enumerate(RIGHT_ORDER, start=60):
        fill_row(ws, idx, 1, left_map[label])
    for idx in range(79, 82):
        for col in range(1, 14):
            ws.cell(idx, col, None)

    # Right detail block Q28:AF50
    headers = [
        "团队",
        "样本数 ",
        "老师数",
        "感染力分数",
        "声音洪亮分数 ",
        "知识点讲错",
        "知识点漏讲",
        "课堂规则提及",
        "开口时长",
        "课中正确率 ",
        "上台次数",
        "学员正面动作次数",
        "读题审题 ",
        "关键提问 ",
        "情绪策略 ",
        "课程预告 ",
    ]
    fill_row(ws, 28, 17, headers)
    for idx, label in enumerate(RIGHT_ORDER, start=29):
        fill_row(ws, idx, 17, right_map[label])
    for idx in range(48, 51):
        for col in range(17, 33):
            ws.cell(idx, col, None)

    # 7月达成情况静态回填，避免公式缓存未刷新时展示空值
    ws7 = wb["7月达成情况"]
    for row_idx in range(4, 23):
        label = ws7.cell(row_idx, 3).value or ws7.cell(row_idx, 2).value
        if not label or label not in top_map or label not in left_map or label not in right_map:
            continue
        top_row = top_map[label]
        left_row = left_map[label]
        right_row = right_map[label]

        teacher_count = top_row[2]
        bi_actual = top_row[1]
        inf_actual = left_row[9]
        interact_actual = left_row[10]
        grooming_values = left_row[3:9]
        grooming_avg = sum(grooming_values) / len(grooming_values) if grooming_values else None

        ws7.cell(row_idx, 4, teacher_count)
        ws7.cell(row_idx, 6, bi_actual)
        ws7.cell(row_idx, 7, None if bi_actual is None else bi_actual - 6)
        ws7.cell(row_idx, 9, inf_actual)
        ws7.cell(row_idx, 10, None if inf_actual is None else inf_actual - 0.5)
        ws7.cell(row_idx, 12, interact_actual)
        ws7.cell(row_idx, 13, "/")
        ws7.cell(row_idx, 14, grooming_avg)

        for col_idx, value in enumerate(grooming_values, start=15):
            ws7.cell(row_idx, col_idx, value)

        ws7.cell(row_idx, 21, left_row[11])
        ws7.cell(row_idx, 22, left_row[12])
        # 7月达成情况 这里只展示从 感染力分数 到 课程预告，不包含 上台次数
        display_metrics = [
            right_row[3],   # 感染力分数
            right_row[4],   # 声音洪亮分数
            right_row[5],   # 知识点讲错
            right_row[6],   # 知识点漏讲
            right_row[7],   # 课堂规则提及
            right_row[8],   # 开口时长
            right_row[9],   # 课中正确率
            right_row[11],  # 学员正面动作次数
            right_row[12],  # 读题审题
            right_row[13],  # 关键提问
            right_row[14],  # 情绪策略
            right_row[15],  # 课程预告
        ]
        for col_idx, value in enumerate(display_metrics, start=23):
            ws7.cell(row_idx, col_idx, value)
        ws7.cell(row_idx, 35, None)

    wb.save(TARGET_PATH)
    print(TARGET_PATH)


if __name__ == "__main__":
    main()
